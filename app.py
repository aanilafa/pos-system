import streamlit as st
import sqlite3
import pandas as pd
from datetime import datetime, date
import io
import random
import time
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

try:
    import requests
    GSHEETS_LIB_AVAILABLE = True
except ImportError:
    GSHEETS_LIB_AVAILABLE = False

# ============================================================
# PHONE REPAIR SHOP - CONFIGURATION
# ============================================================
DB_FILE = "inventory.db"  # keeps the existing app database location
ADMIN_PIN = "1234"         # CHANGE THIS before deploying

GSHEET_WEBAPP_URL = "https://script.google.com/macros/s/AKfycbxtk6yTSHJG6PGy2ZzGA7oczM9bDm78-o-FDuy_tZZfi-Puoltms8KHqgRSt0-26dI/exec"
GSHEET_WEBAPP_SECRET = "POS-SI-2026-9xK7mQ4vT8pL2"
GSHEET_SHARE_URL = "https://docs.google.com/spreadsheets/d/1eokIRdiCSkEIkSSckMT6kMERMJDuA1i7iUIk6OWzqa8/edit?gid=1843822426#gid=1843822426"

st.set_page_config(page_title="Phone Repair Shop System", layout="wide", page_icon="📱")

st.markdown("""
<style>
.block-container { padding-top: 1.2rem; padding-bottom: 2rem; }
div[data-testid="stHorizontalBlock"] button { border-radius: 7px !important; font-weight: 600 !important; min-height: 40px !important; }
.job-card { border:1px solid #e5e7eb; border-radius:10px; padding:14px; margin-bottom:10px; background:#fafafa; }
.status-pill { display:inline-block; padding:4px 9px; border-radius:12px; font-size:12px; font-weight:700; }
.small-note { color:#6b7280; font-size:0.9rem; }
</style>
""", unsafe_allow_html=True)

PAYMENT_METHODS = ["Cash", "Card", "Bank Transfer"]
REPAIR_STATUSES = ["Received", "Diagnosing", "Waiting for Parts", "Repairing", "Ready for Collection", "Collected", "Cancelled"]
DEVICE_TYPES = ["Smartphone", "Tablet", "Smartwatch", "Other"]
PRODUCT_CATEGORIES = ["Screens", "Batteries", "Charging Parts", "Cameras", "Speakers", "Microphones", "Connectors", "IC / Boards", "Tools", "Accessories", "Other"]

GSHEET_HEADERS = [
    "Sale ID", "Timestamp", "Receipt Ref", "Product", "Qty",
    "Unit Cost", "Unit Price", "Discount", "Total", "Profit",
    "Payment Method", "Cashier",
]
GSHEET_INVENTORY_HEADERS = [
    "Product ID", "Product", "Price Tier / Variant", "Category",
    "Unit Cost", "Selling Price", "Margin", "Stock", "Status",
]

# ============================================================
# DATABASE
# ============================================================
def get_connection():
    conn = sqlite3.connect(DB_FILE)
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def init_db():
    conn = get_connection()
    cur = conn.cursor()

    cur.execute("""CREATE TABLE IF NOT EXISTS app_settings (
        key TEXT PRIMARY KEY, value TEXT NOT NULL
    )""")

    # Existing inventory/sales tables are retained for compatibility with the user's app.
    cur.execute("""CREATE TABLE IF NOT EXISTS products (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT,
        variant_label TEXT DEFAULT '',
        category TEXT,
        cost_price REAL DEFAULT 0,
        price REAL DEFAULT 0,
        stock INTEGER DEFAULT 0
    )""")
    cur.execute("""CREATE TABLE IF NOT EXISTS cashiers (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT UNIQUE
    )""")
    cur.execute("""CREATE TABLE IF NOT EXISTS sales (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        receipt_id TEXT,
        product_name TEXT,
        quantity INTEGER,
        unit_cost REAL DEFAULT 0,
        unit_price REAL DEFAULT 0,
        discount_amount REAL DEFAULT 0,
        total_price REAL DEFAULT 0,
        payment_method TEXT,
        cashier TEXT,
        timestamp DATETIME DEFAULT CURRENT_TIMESTAMP
    )""")

    # Phone repair tables.
    cur.execute("""CREATE TABLE IF NOT EXISTS customers (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT NOT NULL,
        phone TEXT DEFAULT '',
        email TEXT DEFAULT '',
        notes TEXT DEFAULT '',
        created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
        updated_at DATETIME DEFAULT CURRENT_TIMESTAMP
    )""")
    cur.execute("""CREATE TABLE IF NOT EXISTS repair_jobs (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        job_ref TEXT UNIQUE NOT NULL,
        customer_id INTEGER,
        customer_name TEXT NOT NULL,
        customer_phone TEXT DEFAULT '',
        device_type TEXT DEFAULT 'Smartphone',
        brand TEXT DEFAULT '',
        model TEXT DEFAULT '',
        imei_serial TEXT DEFAULT '',
        device_condition TEXT DEFAULT '',
        customer_issue TEXT DEFAULT '',
        diagnosis TEXT DEFAULT '',
        repair_notes TEXT DEFAULT '',
        technician TEXT DEFAULT '',
        status TEXT DEFAULT 'Received',
        quoted_amount REAL DEFAULT 0,
        deposit_amount REAL DEFAULT 0,
        paid_amount REAL DEFAULT 0,
        created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
        updated_at DATETIME DEFAULT CURRENT_TIMESTAMP,
        ready_at DATETIME,
        collected_at DATETIME,
        FOREIGN KEY(customer_id) REFERENCES customers(id) ON DELETE SET NULL
    )""")
    cur.execute("""CREATE TABLE IF NOT EXISTS repair_job_parts (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        repair_job_id INTEGER NOT NULL,
        product_id INTEGER,
        product_name TEXT NOT NULL,
        quantity INTEGER DEFAULT 1,
        unit_cost REAL DEFAULT 0,
        unit_price REAL DEFAULT 0,
        FOREIGN KEY(repair_job_id) REFERENCES repair_jobs(id) ON DELETE CASCADE,
        FOREIGN KEY(product_id) REFERENCES products(id) ON DELETE SET NULL
    )""")
    cur.execute("""CREATE TABLE IF NOT EXISTS repair_payments (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        repair_job_id INTEGER NOT NULL,
        amount REAL NOT NULL,
        payment_method TEXT NOT NULL,
        cashier TEXT DEFAULT '',
        reference TEXT DEFAULT '',
        created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY(repair_job_id) REFERENCES repair_jobs(id) ON DELETE CASCADE
    )""")

    # Safe migrations for old databases.
    for table, additions in {
        "products": {
            "variant_label": "TEXT DEFAULT ''", "cost_price": "REAL DEFAULT 0", "price": "REAL DEFAULT 0", "stock": "INTEGER DEFAULT 0"
        },
        "sales": {
            "receipt_id": "TEXT", "unit_cost": "REAL DEFAULT 0", "unit_price": "REAL DEFAULT 0", "discount_amount": "REAL DEFAULT 0"
        },
        "repair_jobs": {
            "device_condition": "TEXT DEFAULT ''", "diagnosis": "TEXT DEFAULT ''", "repair_notes": "TEXT DEFAULT ''",
            "technician": "TEXT DEFAULT ''", "deposit_amount": "REAL DEFAULT 0", "paid_amount": "REAL DEFAULT 0",
            "ready_at": "DATETIME", "collected_at": "DATETIME"
        }
    }.items():
        cur.execute(f"PRAGMA table_info({table})")
        existing = {row[1] for row in cur.fetchall()}
        for col, definition in additions.items():
            if col not in existing:
                cur.execute(f"ALTER TABLE {table} ADD COLUMN {col} {definition}")

    conn.commit()
    conn.close()


init_db()

# ============================================================
# HELPERS
# ============================================================
def money(v):
    return f"${float(v or 0):,.2f}"


def now_text():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def generate_ref(prefix="JOB"):
    return f"{prefix}-{datetime.now().strftime('%y%m%d')}-{random.randint(1000, 9999)}"


def get_setting(key, default=None):
    conn = get_connection()
    try:
        row = conn.execute("SELECT value FROM app_settings WHERE key=?", (key,)).fetchone()
        return row[0] if row else default
    finally:
        conn.close()


def set_setting(key, value):
    conn = get_connection()
    try:
        conn.execute("INSERT INTO app_settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (key, str(value)))
        conn.commit()
    finally:
        conn.close()


def status_label(status):
    return status


def get_cashiers():
    conn = get_connection()
    try:
        return pd.read_sql_query("SELECT * FROM cashiers ORDER BY name", conn)
    finally:
        conn.close()


def gsheet_config(key, fallback, url=False):
    try:
        val = st.secrets.get(key)
    except Exception:
        val = None
    val = val or fallback
    if not val:
        return None
    val = str(val).strip()
    if url and not val.startswith(("https://", "http://")):
        return None
    return val


def gsheet_is_configured():
    return GSHEETS_LIB_AVAILABLE and bool(gsheet_config("gsheet_webapp_url", GSHEET_WEBAPP_URL, True))


def gsheet_url():
    return gsheet_config("gsheet_share_url", GSHEET_SHARE_URL, True)


def post_gsheet(payload, attempts=3):
    if not gsheet_is_configured():
        return {"configured": False, "ok": False, "error": None, "added": 0, "updated": 0}
    secret = gsheet_config("gsheet_webapp_secret", GSHEET_WEBAPP_SECRET) or ""
    url = gsheet_config("gsheet_webapp_url", GSHEET_WEBAPP_URL, True)
    payload = dict(payload)
    payload["secret"] = secret
    last_error = "Unknown Google Sheets error."
    for attempt in range(1, attempts + 1):
        try:
            response = requests.post(url, json=payload, timeout=20)
            if response.status_code != 200:
                last_error = f"HTTP {response.status_code}: {response.text[:250]}"
            else:
                try:
                    data = response.json()
                except ValueError:
                    last_error = f"Web App did not return JSON: {response.text[:250]}"
                else:
                    if data.get("status") == "ok":
                        data.update({"configured": True, "attempts": attempt})
                        return data
                    last_error = data.get("message", "Unknown Web App error.")
        except requests.exceptions.RequestException as exc:
            last_error = str(exc)
        if attempt < attempts:
            time.sleep(2 ** (attempt - 1))
    return {"configured": True, "ok": False, "attempts": attempts, "error": last_error, "added": 0, "updated": 0}


def sync_sales_to_gsheet(df):
    if df is None or df.empty:
        return {"configured": True, "ok": True, "added": 0}
    rows = []
    for _, r in df.iterrows():
        qty = int(r.get("quantity", 0) or 0)
        cost = float(r.get("unit_cost", 0) or 0)
        total = float(r.get("total_price", 0) or 0)
        rows.append([
            str(int(r.get("id", 0))), str(r.get("timestamp", "")), str(r.get("receipt_id", "")),
            str(r.get("product_name", "")), qty, cost, float(r.get("unit_price", 0) or 0),
            float(r.get("discount_amount", 0) or 0), total, total - cost * qty,
            str(r.get("payment_method", "")), str(r.get("cashier", ""))
        ])
    return post_gsheet({"headers": GSHEET_HEADERS, "rows": rows})


def sync_inventory_to_gsheet():
    conn = get_connection()
    try:
        df = pd.read_sql_query("SELECT * FROM products ORDER BY id", conn)
    finally:
        conn.close()
    rows = []
    for _, r in df.iterrows():
        stock = int(r.get("stock", 0) or 0)
        status = "Out of Stock" if stock <= 0 else ("Low Stock" if stock <= 3 else "In Stock")
        cost = float(r.get("cost_price", 0) or 0)
        price = float(r.get("price", 0) or 0)
        rows.append([str(int(r["id"])), str(r.get("name", "")), str(r.get("variant_label", "") or ""), str(r.get("category", "")), cost, price, price-cost, stock, status])
    return post_gsheet({"action": "sync_inventory", "headers": GSHEET_INVENTORY_HEADERS, "rows": rows})


def clear_google(action):
    return post_gsheet({"action": action})


def create_receipt_docx(title, ref, customer, device, items, total, paid, balance, cashier, payment_method):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.5); section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5); section.right_margin = Inches(0.5)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PHONE REPAIR SHOP\n"); r.bold = True; r.font.size = Pt(18)
    r = p.add_run(title); r.font.size = Pt(11)
    doc.add_paragraph("-" * 90)
    p = doc.add_paragraph()
    p.add_run("Reference: ").bold = True; p.add_run(f"{ref}\n")
    p.add_run("Date & Time: ").bold = True; p.add_run(f"{now_text()}\n")
    p.add_run("Customer: ").bold = True; p.add_run(f"{customer}\n")
    p.add_run("Device: ").bold = True; p.add_run(f"{device}\n")
    p.add_run("Cashier/Technician: ").bold = True; p.add_run(f"{cashier}\n")
    p.add_run("Payment: ").bold = True; p.add_run(f"{payment_method}\n")
    table = doc.add_table(rows=1, cols=3)
    for i, h in enumerate(["Description", "Qty", "Amount"]):
        table.rows[0].cells[i].text = h
    for name, qty, amount in items:
        row = table.add_row().cells
        row[0].text = str(name); row[1].text = str(qty); row[2].text = money(amount)
    doc.add_paragraph("-" * 90)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run(f"TOTAL: {money(total)}\n").bold = True
    p.add_run(f"PAID: {money(paid)}\n")
    p.add_run(f"BALANCE: {money(balance)}\n").bold = True
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Thank you for choosing our repair service!").italic = True
    out = io.BytesIO(); doc.save(out); return out.getvalue()


def export_excel(df, sheet_name="Report"):
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name[:31])
        ws = writer.sheets[sheet_name[:31]]
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        thin = Border(left=Side(style="thin", color="D9D9D9"), right=Side(style="thin", color="D9D9D9"), top=Side(style="thin", color="D9D9D9"), bottom=Side(style="thin", color="D9D9D9"))
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(1, c); cell.font = header_font; cell.fill = header_fill; cell.alignment = Alignment(horizontal="center"); cell.border = thin
            width = max(len(str(ws.cell(1,c).value or "")), *(len(str(ws.cell(r,c).value or "")) for r in range(2, min(ws.max_row, 200)+1))) + 3
            ws.column_dimensions[get_column_letter(c)].width = min(max(width, 12), 35)
            for r in range(2, ws.max_row+1): ws.cell(r,c).border = thin
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    return buf.getvalue()

# ============================================================
# DAILY CASH SALES RESET - REPAIR JOBS NEVER RESET
# ============================================================
def daily_sales_reset():
    key = "phone_shop_last_sales_day"
    today = date.today().isoformat()
    last = get_setting(key)
    if last and last != today:
        conn = get_connection()
        try:
            conn.execute("DELETE FROM sales")
            conn.execute("DELETE FROM sqlite_sequence WHERE name='sales'")
            conn.commit()
        finally:
            conn.close()
    if last != today:
        set_setting(key, today)

daily_sales_reset()

# ============================================================
# SESSION STATE
# ============================================================
DEFAULTS = {
    "cashier": "", "last_receipt": None, "job_receipt": None,
    "admin_authenticated": False, "repair_filter": "All"
}
for key, value in DEFAULTS.items():
    if key not in st.session_state:
        st.session_state[key] = value

# ============================================================
# SIDEBAR
# ============================================================
st.sidebar.markdown("### 📱 Phone Repair Shop")
st.sidebar.caption("Repair jobs • Parts inventory • Daily payments")
role = st.sidebar.radio("Navigation", [
    "🔧 Repair Jobs", "💳 Sales / Payments", "📦 Parts Inventory", "📊 Reports", "🔐 Admin"
], label_visibility="collapsed")

cashiers_df = get_cashiers()
cashier_names = cashiers_df["name"].tolist() if not cashiers_df.empty else []
if role != "🔐 Admin":
    if cashier_names:
        idx = cashier_names.index(st.session_state.cashier) if st.session_state.cashier in cashier_names else 0
        st.sidebar.selectbox("Staff", cashier_names, index=idx, key="cashier")
    else:
        st.sidebar.text_input("Staff Name", key="cashier", value="Staff")

# ============================================================
# REPAIR JOBS
# ============================================================
if role == "🔧 Repair Jobs":
    st.title("🔧 Repair Jobs")
    st.caption("Receive phones, track diagnosis and repair progress, record deposits/payments, and close jobs when collected.")

    conn = get_connection()
    jobs_count = conn.execute("SELECT COUNT(*) FROM repair_jobs WHERE status NOT IN ('Collected','Cancelled')").fetchone()[0]
    ready_count = conn.execute("SELECT COUNT(*) FROM repair_jobs WHERE status='Ready for Collection'").fetchone()[0]
    balance_total = conn.execute("SELECT COALESCE(SUM(quoted_amount-paid_amount),0) FROM repair_jobs WHERE status NOT IN ('Collected','Cancelled')").fetchone()[0]
    conn.close()
    m1,m2,m3 = st.columns(3)
    m1.metric("Open Jobs", jobs_count); m2.metric("Ready for Collection", ready_count); m3.metric("Outstanding Balance", money(balance_total))

    tab_new, tab_jobs = st.tabs(["➕ New Repair Job", "📋 Repair Job List"])
    with tab_new:
        st.markdown("#### Customer & Device")
        c1,c2 = st.columns(2)
        with c1:
            customer_name = st.text_input("Customer Name *", key="new_customer_name")
            customer_phone = st.text_input("Customer Phone", key="new_customer_phone")
            customer_email = st.text_input("Customer Email", key="new_customer_email")
        with c2:
            device_type = st.selectbox("Device Type", DEVICE_TYPES, key="new_device_type")
            brand = st.text_input("Brand", placeholder="Apple, Samsung, Tecno, Infinix...", key="new_brand")
            model = st.text_input("Model", key="new_model")
        imei = st.text_input("IMEI / Serial Number", key="new_imei")
        condition = st.text_area("Device Condition / Accessories", placeholder="Cracked glass, scratches, charger left with device, etc.", key="new_condition")
        issue = st.text_area("Customer's Reported Problem *", placeholder="Phone won't charge, broken screen, no power...", key="new_issue")
        quoted = st.number_input("Quoted Repair Price ($)", min_value=0.0, step=1.0, format="%.2f", key="new_quoted")
        deposit = st.number_input("Deposit / Initial Payment ($)", min_value=0.0, max_value=float(quoted), step=1.0, format="%.2f", key="new_deposit")
        tech = st.text_input("Technician", value=st.session_state.get("cashier", ""), key="new_technician")
        notes = st.text_area("Initial Notes", key="new_notes")

        if st.button("Create Repair Job", type="primary", use_container_width=True):
            if not customer_name.strip() or not issue.strip():
                st.error("Customer name and reported problem are required.")
            else:
                conn = get_connection()
                try:
                    cur = conn.cursor()
                    cur.execute("INSERT INTO customers(name,phone,email,notes) VALUES(?,?,?,?)", (customer_name.strip(), customer_phone.strip(), customer_email.strip(), ""))
                    customer_id = cur.lastrowid
                    ref = generate_ref("JOB")
                    cur.execute("""INSERT INTO repair_jobs(job_ref,customer_id,customer_name,customer_phone,device_type,brand,model,imei_serial,device_condition,customer_issue,technician,status,quoted_amount,deposit_amount,paid_amount,repair_notes)
                        VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", (ref,customer_id,customer_name.strip(),customer_phone.strip(),device_type,brand.strip(),model.strip(),imei.strip(),condition.strip(),issue.strip(),tech.strip(),"Received",quoted,deposit,deposit,notes.strip()))
                    job_id = cur.lastrowid
                    if deposit > 0:
                        cur.execute("INSERT INTO repair_payments(repair_job_id,amount,payment_method,cashier,reference) VALUES(?,?,?,?,?)", (job_id,deposit,"Cash",st.session_state.cashier,ref))
                        receipt_id = f"PAY-{random.randint(100000,999999)}"
                        cur.execute("INSERT INTO sales(receipt_id,product_name,quantity,unit_cost,unit_price,discount_amount,total_price,payment_method,cashier) VALUES(?,?,?,?,?,?,?,?,?)", (receipt_id,f"Repair Deposit - {ref}",1,0,deposit,0,deposit,"Cash",st.session_state.cashier))
                    conn.commit()
                    st.session_state.job_receipt = ref
                finally:
                    conn.close()
                st.success(f"Repair job {ref} created successfully.")
                st.rerun()

    with tab_jobs:
        f1,f2,f3 = st.columns([2,1,1])
        search = f1.text_input("Search", placeholder="Job ref, customer, phone, IMEI, model...")
        selected_status = f2.selectbox("Status", ["All"] + REPAIR_STATUSES)
        sort_order = f3.selectbox("Order", ["Newest First", "Oldest First"])
        conn = get_connection()
        query = "SELECT * FROM repair_jobs WHERE 1=1"
        params=[]
        if selected_status != "All": query += " AND status=?"; params.append(selected_status)
        if search:
            query += " AND (job_ref LIKE ? OR customer_name LIKE ? OR customer_phone LIKE ? OR imei_serial LIKE ? OR model LIKE ? OR brand LIKE ?)"
            q=f"%{search}%"; params += [q,q,q,q,q,q]
        query += " ORDER BY created_at " + ("DESC" if sort_order == "Newest First" else "ASC")
        jobs = pd.read_sql_query(query, conn, params=params)
        conn.close()

        if jobs.empty:
            st.info("No repair jobs found.")
        else:
            st.dataframe(jobs[["job_ref","customer_name","customer_phone","brand","model","status","quoted_amount","paid_amount","created_at"]].rename(columns={"job_ref":"Job Ref","customer_name":"Customer","customer_phone":"Phone","brand":"Brand","model":"Model","status":"Status","quoted_amount":"Quoted","paid_amount":"Paid","created_at":"Created"}), use_container_width=True, hide_index=True)
            st.download_button("⬇️ Export Repair Jobs", export_excel(jobs, "Repair Jobs"), "repair_jobs.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

            st.markdown("#### Manage Job")
            ref = st.selectbox("Select Job", jobs["job_ref"].tolist())
            conn = get_connection(); job = conn.execute("SELECT * FROM repair_jobs WHERE job_ref=?", (ref,)).fetchone(); cols=[d[0] for d in conn.execute("SELECT * FROM repair_jobs LIMIT 1").description]; conn.close()
            jobd=dict(zip(cols,job))
            remaining=max(0,float(jobd["quoted_amount"] or 0)-float(jobd["paid_amount"] or 0))
            st.info(f"**{jobd['job_ref']}** — {jobd['customer_name']} — {jobd['brand']} {jobd['model']} — Balance: **{money(remaining)}**")
            a,b,c = st.columns(3)
            with a:
                new_status=st.selectbox("Status", REPAIR_STATUSES, index=REPAIR_STATUSES.index(jobd["status"]) if jobd["status"] in REPAIR_STATUSES else 0)
            with b:
                diagnosis=st.text_area("Diagnosis", value=jobd["diagnosis"] or "")
            with c:
                technician=st.text_input("Technician", value=jobd["technician"] or "")
            repair_notes=st.text_area("Repair Notes", value=jobd["repair_notes"] or "")
            quoted_edit=st.number_input("Quoted Amount ($)", min_value=0.0, value=float(jobd["quoted_amount"] or 0), step=1.0)
            if st.button("Save Job Changes", type="primary"):
                conn=get_connection(); conn.execute("UPDATE repair_jobs SET status=?,diagnosis=?,technician=?,repair_notes=?,quoted_amount=?,updated_at=CURRENT_TIMESTAMP,ready_at=CASE WHEN ?='Ready for Collection' THEN CURRENT_TIMESTAMP ELSE ready_at END,collected_at=CASE WHEN ?='Collected' THEN CURRENT_TIMESTAMP ELSE collected_at END WHERE id=?", (new_status,diagnosis,technician,repair_notes,quoted_edit,new_status,new_status,jobd["id"])); conn.commit(); conn.close(); st.success("Repair job updated."); st.rerun()

            st.markdown("##### Record Payment")
            p1,p2,p3=st.columns(3)
            with p1: payment_amount=st.number_input("Payment Amount ($)", min_value=0.0, max_value=remaining, step=1.0, format="%.2f", key="job_payment_amount")
            with p2: payment_method=st.selectbox("Payment Method", PAYMENT_METHODS, key="job_payment_method")
            with p3: payment_ref=st.text_input("Payment Reference", key="job_payment_ref")
            if st.button("Record Payment", use_container_width=True):
                if payment_amount <= 0:
                    st.error("Enter a payment amount.")
                else:
                    conn=get_connection()
                    try:
                        cur=conn.cursor(); cur.execute("INSERT INTO repair_payments(repair_job_id,amount,payment_method,cashier,reference) VALUES(?,?,?,?,?)", (jobd["id"],payment_amount,payment_method,st.session_state.cashier,payment_ref.strip()))
                        cur.execute("UPDATE repair_jobs SET paid_amount=paid_amount+?,deposit_amount=deposit_amount+?,updated_at=CURRENT_TIMESTAMP WHERE id=?", (payment_amount,payment_amount,jobd["id"]))
                        receipt_id=f"PAY-{random.randint(100000,999999)}"
                        cur.execute("INSERT INTO sales(receipt_id,product_name,quantity,unit_cost,unit_price,discount_amount,total_price,payment_method,cashier) VALUES(?,?,?,?,?,?,?,?,?)", (receipt_id,f"Repair Payment - {jobd['job_ref']}",1,0,payment_amount,0,payment_amount,payment_method,st.session_state.cashier))
                        conn.commit()
                    finally: conn.close()
                    st.success(f"Payment of {money(payment_amount)} recorded."); st.rerun()

            conn=get_connection(); parts=pd.read_sql_query("SELECT * FROM repair_job_parts WHERE repair_job_id=?", conn, params=(jobd["id"],)); conn.close()
            if not parts.empty:
                st.markdown("##### Parts Used")
                st.dataframe(parts[["product_name","quantity","unit_price"]], use_container_width=True, hide_index=True)
            st.markdown("##### Job History / Payments")
            conn=get_connection(); payments=pd.read_sql_query("SELECT created_at,amount,payment_method,cashier,reference FROM repair_payments WHERE repair_job_id=? ORDER BY created_at DESC", conn, params=(jobd["id"],)); conn.close()
            if payments.empty: st.caption("No payments recorded.")
            else: st.dataframe(payments, use_container_width=True, hide_index=True)

# ============================================================
# SALES / PAYMENTS
# ============================================================
elif role == "💳 Sales / Payments":
    st.title("💳 Daily Sales & Payments")
    st.caption("This screen records daily cash/card/bank transactions. Repair jobs remain stored separately and do not reset.")
    conn=get_connection(); sales=pd.read_sql_query("SELECT * FROM sales ORDER BY timestamp DESC", conn); conn.close()
    today_df=sales[pd.to_datetime(sales["timestamp"]).dt.date == date.today()] if not sales.empty else pd.DataFrame()
    c1,c2,c3=st.columns(3); c1.metric("Today's Revenue", money(today_df["total_price"].sum() if not today_df.empty else 0)); c2.metric("Transactions", len(today_df)); c3.metric("Cash", money(today_df.loc[today_df.payment_method=="Cash","total_price"].sum() if not today_df.empty else 0))
    if sales.empty: st.info("No payments recorded today.")
    else:
        st.dataframe(sales.rename(columns={"receipt_id":"Reference","product_name":"Description","quantity":"Qty","unit_price":"Amount","total_price":"Total","payment_method":"Payment","cashier":"Staff","timestamp":"Time"})[["id","Reference","Description","Qty","Amount","Total","Payment","Staff","Time"]], use_container_width=True, hide_index=True)
        st.download_button("⬇️ Export Daily Payments", export_excel(today_df if not today_df.empty else sales, "Daily Payments"), "daily_payments.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

# ============================================================
# PARTS INVENTORY
# ============================================================
elif role == "📦 Parts Inventory":
    st.title("📦 Parts & Accessories Inventory")
    st.caption("Manage replacement screens, batteries, charging parts, tools and other stock used by the repair shop.")
    if gsheet_is_configured():
        g1,g2=st.columns(2)
        with g1:
            if st.button("🔄 Sync Inventory to Google Sheets", use_container_width=True):
                result=sync_inventory_to_gsheet(); st.success(f"Inventory synced ({result.get('updated',0)} rows).") if result.get("ok") else st.error(result.get("error","Sync failed"))
        with g2:
            link=gsheet_url()
            if link: st.link_button("🔗 Open Google Sheet", link, use_container_width=True)
    list_tab, add_tab = st.tabs(["📋 Stock List", "➕ Add / Restock Part"])
    with list_tab:
        conn=get_connection(); products=pd.read_sql_query("SELECT * FROM products ORDER BY name, id", conn); conn.close()
        if products.empty: st.info("No parts in inventory yet.")
        else:
            products["Status"]=products.stock.apply(lambda x:"Out of Stock" if x<=0 else ("Low Stock" if x<=3 else "In Stock"))
            products["Margin"]=products.price-products.cost_price
            st.dataframe(products[["id","name","variant_label","category","cost_price","price","Margin","stock","Status"]].rename(columns={"id":"ID","name":"Part","variant_label":"Variant","category":"Category","cost_price":"Cost","price":"Selling Price","stock":"Stock"}), use_container_width=True, hide_index=True)
            st.download_button("⬇️ Export Inventory", export_excel(products, "Inventory"), "phone_parts_inventory.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            st.markdown("##### Edit Stock / Price")
            pid=st.selectbox("Part", products.id.tolist(), format_func=lambda x: f"#{x} — {products.loc[products.id==x,'name'].iloc[0]}")
            row=products[products.id==pid].iloc[0]
            e1,e2,e3=st.columns(3)
            with e1: new_stock=st.number_input("Stock", min_value=0, value=int(row.stock), step=1)
            with e2: new_cost=st.number_input("Cost ($)", min_value=0.0, value=float(row.cost_price), step=1.0)
            with e3: new_price=st.number_input("Selling Price ($)", min_value=0.0, value=float(row.price), step=1.0)
            if st.button("Save Inventory Changes", type="primary"):
                conn=get_connection(); conn.execute("UPDATE products SET stock=?,cost_price=?,price=? WHERE id=?", (new_stock,new_cost,new_price,pid)); conn.commit(); conn.close();
                if gsheet_is_configured(): sync_inventory_to_gsheet()
                st.success("Inventory updated."); st.rerun()
    with add_tab:
        st.text_input("Part Name", key="part_name")
        st.text_input("Variant / Compatible Model", placeholder="iPhone 11 / A2111", key="part_variant")
        st.selectbox("Category", PRODUCT_CATEGORIES, key="part_category")
        a,b,c=st.columns(3)
        with a: cost=st.number_input("Unit Cost ($)", min_value=0.0, step=1.0, key="part_cost")
        with b: markup=st.number_input("Markup ($)", min_value=0.0, step=1.0, key="part_markup")
        with c: st.metric("Selling Price", money(cost+markup))
        qty=st.number_input("Quantity to Add", min_value=0, step=1, key="part_qty")
        if st.button("Save / Restock Part", type="primary", use_container_width=True):
            name=st.session_state.part_name.strip(); variant=st.session_state.part_variant.strip(); price=cost+markup
            if not name: st.error("Part name is required.")
            else:
                conn=get_connection(); existing=conn.execute("SELECT id FROM products WHERE name=? AND IFNULL(variant_label,'')=? AND price=?", (name,variant,price)).fetchone()
                if existing: conn.execute("UPDATE products SET stock=stock+?,cost_price=?,category=? WHERE id=?", (qty,cost,st.session_state.part_category,existing[0]))
                else: conn.execute("INSERT INTO products(name,variant_label,category,cost_price,price,stock) VALUES(?,?,?,?,?,?)", (name,variant,st.session_state.part_category,cost,price,qty))
                conn.commit(); conn.close()
                if gsheet_is_configured(): sync_inventory_to_gsheet()
                st.success("Part saved/restocked."); st.rerun()

# ============================================================
# REPORTS
# ============================================================
elif role == "📊 Reports":
    st.title("📊 Phone Repair Reports")
    tab1,tab2,tab3=st.tabs(["Repair Jobs","Revenue","Customers"])
    with tab1:
        conn=get_connection(); jobs=pd.read_sql_query("SELECT * FROM repair_jobs ORDER BY created_at DESC", conn); conn.close()
        if jobs.empty: st.info("No repair jobs yet.")
        else:
            r1,r2,r3,r4=st.columns(4)
            r1.metric("Total Jobs",len(jobs)); r2.metric("Open",int((~jobs.status.isin(["Collected","Cancelled"])).sum())); r3.metric("Collected",int((jobs.status=="Collected").sum())); r4.metric("Outstanding",money((jobs.quoted_amount-jobs.paid_amount).clip(lower=0).sum()))
            st.dataframe(jobs.groupby("status").size().reset_index(name="Jobs"), use_container_width=True, hide_index=True)
            st.download_button("⬇️ Export All Repair Jobs", export_excel(jobs,"Repair Jobs"),"repair_jobs_report.xlsx","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    with tab2:
        conn=get_connection(); sales=pd.read_sql_query("SELECT * FROM sales ORDER BY timestamp DESC", conn); conn.close()
        if sales.empty: st.info("No revenue records.")
        else:
            sales["Date"]=pd.to_datetime(sales.timestamp).dt.date
            summary=sales.groupby(["Date","payment_method"],as_index=False)["total_price"].sum().rename(columns={"payment_method":"Payment Method","total_price":"Revenue"})
            st.dataframe(summary,use_container_width=True,hide_index=True)
            st.download_button("⬇️ Export Revenue Report",export_excel(sales,"Revenue"),"revenue_report.xlsx","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    with tab3:
        conn=get_connection(); customers=pd.read_sql_query("SELECT * FROM customers ORDER BY updated_at DESC", conn); conn.close()
        st.dataframe(customers,use_container_width=True,hide_index=True) if not customers.empty else st.info("No customers yet.")

# ============================================================
# ADMIN
# ============================================================
elif role == "🔐 Admin":
    st.title("🔐 Admin Dashboard")
    if not st.session_state.admin_authenticated:
        st.info("Enter the Admin PIN to access management tools.")
        pin=st.text_input("Admin PIN", type="password")
        if st.button("Unlock Admin", type="primary"):
            if pin == ADMIN_PIN:
                st.session_state.admin_authenticated=True; st.rerun()
            else: st.error("Incorrect Admin PIN.")
    else:
        if st.button("🔒 Lock Admin"):
            st.session_state.admin_authenticated=False; st.rerun()
        tabs=st.tabs(["👥 Staff","🧹 Data Management","☁️ Google Sheets","⚙️ System"])
        with tabs[0]:
            st.subheader("Staff / Cashiers")
            name=st.text_input("Staff Name", key="admin_staff_name")
            if st.button("Add Staff", type="primary"):
                if name.strip():
                    try:
                        conn=get_connection(); conn.execute("INSERT INTO cashiers(name) VALUES(?)",(name.strip(),)); conn.commit(); conn.close(); st.success("Staff member added."); st.rerun()
                    except sqlite3.IntegrityError: st.error("That staff member already exists.")
                else: st.error("Enter a name.")
            conn=get_connection(); staff=pd.read_sql_query("SELECT * FROM cashiers ORDER BY name",conn); conn.close(); st.dataframe(staff,use_container_width=True,hide_index=True) if not staff.empty else st.caption("No staff registered.")
        with tabs[1]:
            st.subheader("Clear Data")
            st.warning("These actions permanently delete local records. Repair jobs and customer records are separate from daily sales.")
            if st.button("Clear Daily Sales / Payments", use_container_width=True):
                conn=get_connection(); conn.execute("DELETE FROM sales"); conn.execute("DELETE FROM sqlite_sequence WHERE name='sales'"); conn.commit(); conn.close();
                if gsheet_is_configured(): clear_google("clear_sales")
                st.success("Daily sales/payment records cleared.")
            if st.button("Clear Parts Inventory", use_container_width=True):
                conn=get_connection(); conn.execute("DELETE FROM products"); conn.execute("DELETE FROM sqlite_sequence WHERE name='products'"); conn.commit(); conn.close();
                if gsheet_is_configured(): clear_google("clear_inventory")
                st.success("Parts inventory cleared.")
            st.markdown("#### Full Reset")
            confirm=st.text_input('Type "RESET PHONE SHOP" to confirm', key="full_reset_confirm")
            if st.button("⚠️ Reset Repair Shop Data", type="secondary"):
                if confirm == "RESET PHONE SHOP":
                    conn=get_connection()
                    try:
                        for table in ["repair_payments","repair_job_parts","repair_jobs","customers","sales","products","cashiers"]: conn.execute(f"DELETE FROM {table}")
                        conn.commit()
                    finally: conn.close()
                    if gsheet_is_configured(): clear_google("clear_all")
                    st.success("Phone repair shop data has been reset."); st.rerun()
                else: st.error("Confirmation text does not match.")
        with tabs[2]:
            st.subheader("Google Sheets")
            if not gsheet_is_configured(): st.info("Google Sheets is not configured or the requests package is unavailable.")
            else:
                st.success("Google Sheets connection is configured.")
                if gsheet_url(): st.link_button("🔗 Open Shared Google Sheet", gsheet_url())
                if st.button("🔄 Sync All Daily Sales"):
                    conn=get_connection(); df=pd.read_sql_query("SELECT * FROM sales ORDER BY id",conn); conn.close(); result=sync_sales_to_gsheet(df); st.success(f"Synced {result.get('added',0)} row(s).") if result.get("ok") else st.error(result.get("error","Sync failed"))
                if st.button("🔄 Sync Parts Inventory"):
                    result=sync_inventory_to_gsheet(); st.success(f"Inventory sync complete: {result.get('updated',0)} row(s).") if result.get("ok") else st.error(result.get("error","Sync failed"))
        with tabs[3]:
            st.subheader("System Information")
            st.write("Database:", DB_FILE)
            st.write("Daily sales reset:", "Enabled")
            st.write("Repair jobs reset daily:", "No")
            st.write("Google Sheets:", "Configured" if gsheet_is_configured() else "Not configured")
            st.caption("Change ADMIN_PIN near the top of app.py before deploying.")
